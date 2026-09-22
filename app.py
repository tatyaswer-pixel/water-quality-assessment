from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.utils import secure_filename
from collections import defaultdict
import os
import pandas as pd
from datetime import datetime, timedelta
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
import math
import hashlib
import json
import io
import urllib.parse

from config import config
from models import db, User, WaterQualityData, EvaluationResult, init_db
from fuzzy_logic import (
    WaterQualityFuzzySystem, FUZZY_VERSION,
    get_fuzzy_system, BUILDING_TYPES, DEFAULT_BUILDING_TYPE,
    BUILDING_STANDARDS, WQI_BANDS, score_to_wqi_band, score_to_wqi_message,
    get_param_zones,
)
from xlsx_io import (
    normalize_hospital_xlsx, build_hospital_xlsx_report,
    build_single_result_xlsx, CL2_OZONE_SUBSTITUTE
)
from types import SimpleNamespace
from flask import make_response
REQUIRED_FUZZY_COLS = ['pH', 'BOD', 'COD', 'TSS', 'TDS', 'O&G', 'TKN', 'Sulfide', 'TCB', 'FCB', 'Cl2']

app = Flask(__name__)
app.config.from_object(config[os.environ.get('FLASK_ENV', 'development')])
init_db(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'
login_manager.login_message = 'กรุณาเข้าสู่ระบบก่อนใช้งาน'

fuzzy_system = get_fuzzy_system(DEFAULT_BUILDING_TYPE)  # legacy default (ก) for backward compat

# =========================
# HELPERS
# =========================
def is_missing(x):
    if x is None: return True
    if isinstance(x, str) and x.strip() == '': return True
    try: return pd.isna(x)
    except: return False

def safe_float(x):
    try:
        if x is None: return None
        if isinstance(x, str):
            x = x.strip()
            if x == '': return None
        val = float(x)
        if math.isnan(val) or math.isinf(val): return None
        return val
    except: return None

def to_object(data):
    if isinstance(data, dict): return SimpleNamespace(**{k: to_object(v) for k, v in data.items()})
    if isinstance(data, list): return [to_object(item) for item in data]
    return data

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']


def _ascii_filename(name: str) -> str:
    """แปลงชื่อไฟล์ให้เป็น ASCII ล้วน (ใช้เป็น fallback filename ใน Content-Disposition)

    HTTP headers ต้อง encode ได้ด้วย latin-1 (RFC 7230) ดังนั้น filename parameter
    ต้องเป็น ASCII เท่านั้น. ตัวอักษรที่ไม่ใช่ ASCII จะถูกแทนด้วย transliteration
    (ถ้ามี) หรือ '_'. UTF-8 version จะใช้ใน filename*= parameter แยกต่างหาก
    """
    # Common Thai hospital terms → ASCII
    THAI_MAP = {
        'รพ.': 'RP.', 'รพ': 'RP',
        'โรงพยาบาล': 'Hospital_',
        'สมเด็จ': 'Somdet',
        'มกราคม':'Jan','กุมภาพันธ์':'Feb','มีนาคม':'Mar','เมษายน':'Apr',
        'พฤษภาคม':'May','มิถุนายน':'Jun','กรกฎาคม':'Jul','สิงหาคม':'Aug',
        'กันยายน':'Sep','ตุลาคม':'Oct','พฤศจิกายน':'Nov','ธันวาคม':'Dec',
    }
    s = name
    for th, en in THAI_MAP.items():
        s = s.replace(th, en)
    # Any remaining non-ASCII → _
    out = []
    for ch in s:
        if ord(ch) < 128:
            out.append(ch)
        else:
            out.append('_')
    import re as _re
    s = ''.join(out)
    s = _re.sub(r'_+', '_', s)              # collapse runs of _
    s = s.replace(' ', '_').replace('/', '-')
    s = s.strip('_')
    return s or 'export.xlsx'


def _xlsx_download_response(xlsx_bytes: bytes, filename_thai: str):
    """สร้าง Flask response สำหรับ download ไฟล์ .xlsx อย่างปลอดภัย
       - filename="..." ใช้ ASCII-safe fallback (RFC 7230 compliant)
       - filename*=UTF-8''... ใช้ชื่อไทย URL-encoded (RFC 5987)
       - ใส่ no-cache headers
    """
    from flask import send_file
    fa_ascii = _ascii_filename(filename_thai)
    resp = send_file(
        io.BytesIO(xlsx_bytes),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=fa_ascii,         # <- ASCII only
        max_age=0,
    )
    # สร้าง Content-Disposition ที่ latin-1 encode ได้แน่ ๆ
    cd = f'attachment; filename="{fa_ascii}"; filename*=UTF-8\'\'{urllib.parse.quote(filename_thai)}'
    # sanity check: header ทั้งหมดต้อง encode เป็น latin-1 ได้
    try:
        cd.encode('latin-1')
    except UnicodeEncodeError:
        # fallback 100%: ใช้แค่ ASCII filename อย่างเดียว
        cd = f'attachment; filename="{fa_ascii}"'
    resp.headers['Content-Disposition'] = cd
    resp.headers['Content-Length'] = str(len(xlsx_bytes))
    resp.headers['Cache-Control']  = 'no-cache, no-store, must-revalidate'
    resp.headers['Pragma']         = 'no-cache'
    resp.headers['Expires']        = '0'
    resp.headers['X-Content-Type-Options'] = 'nosniff'
    return resp

def get_float(form, name):
    value = form.get(name)
    return float(value) if value not in (None, '') else None

def normalize_key(k):
    return (k or "").replace('&','').replace(' ','').replace('_','').lower().strip()

def score_to_overall_status(score):
    """แปลงคะแนน 0-100 → Pass / Near Limit / Fail (สอดคล้องกับมาตรฐาน WQI 5 ระดับ)
    - Pass        : score ≥ 60   (Excellent 80-100, Good 60-79)
    - Near Limit  : 40 ≤ score < 60 (Fair 40-59 — เฝ้าระวัง ใกล้เกินมาตรฐาน)
    - Fail        : score < 40   (Poor 20-39 + Very Poor 0-19)
    """
    if score is None: return 'No Data'
    try: s = float(score)
    except: return 'No Data'
    if s >= 60: return 'Pass'
    elif s >= 40: return 'Near Limit'
    return 'Fail'

# =========================
# WATER STANDARDS (per building type)
# =========================
PARAM_DISPLAY_NAMES = {
    'pH':'[pH] ค่าความเป็นกรด-ด่าง','BOD':'[BOD] ความต้องการออกซิเจนทางชีวภาพ',
    'COD':'[COD] ความต้องการออกซิเจนทางเคมี','TSS':'[TSS] ของแข็งแขวนลอยทั้งหมด',
    'TDS':'[TDS] ของแข็งละลายทั้งหมด','O&G':'[O&G] น้ำมันและไขมัน','TKN':'[TKN] ไนโตรเจนรวม',
    'Sulfide':'[SULFIDE] ซัลไฟด์','TCB':'[TCB] โคลิฟอร์มทั้งหมด','FCB':'[FCB] โคลิฟอร์มอุจจาระ','Cl2':'[Cl2] คลอรีนตกค้าง'
}
PARAM_UNITS = {
    'pH':'','BOD':'mg/L','COD':'mg/L','TSS':'mg/L','TDS':'mg/L',
    'O&G':'mg/L','TKN':'mg/L','Sulfide':'mg/L','TCB':'MPN/100mL','FCB':'MPN/100mL','Cl2':'mg/L'
}

def _build_water_standards_for(building_type):
    """สร้าง dict มาตรฐานน้ำสำหรับอาคารแต่ละประเภท (ก/ข/ค/ง)"""
    bt_stds = BUILDING_STANDARDS.get(building_type, BUILDING_STANDARDS[DEFAULT_BUILDING_TYPE])
    out = {}
    for param_name, standard in bt_stds.items():
        if standard is None:
            # ไม่กำหนด — แสดงเป็น "ไม่กำหนด" และไม่มีค่า min/max
            out[param_name] = {
                'name': PARAM_DISPLAY_NAMES.get(param_name, param_name),
                'unit': PARAM_UNITS.get(param_name, ''),
                'type': None, 'min': None, 'max': None,
                'regulated': False,
            }
            continue
        if standard.get('min') is not None and standard.get('max') is not None: std_type = 'range'
        elif standard.get('max') is not None: std_type = 'max'
        else: std_type = None
        out[param_name] = {
            'name': PARAM_DISPLAY_NAMES.get(param_name, param_name),
            'unit': PARAM_UNITS.get(param_name, ''),
            'type': std_type,
            'min': standard.get('min'),
            'max': standard.get('max'),
            'regulated': True,
        }
    return out

WATER_STANDARDS_BY_TYPE = {bt: _build_water_standards_for(bt) for bt in BUILDING_TYPES}
# default (ก) สำหรับ template/route ที่ยังไม่มี context ของ building_type
WATER_STANDARDS = WATER_STANDARDS_BY_TYPE[DEFAULT_BUILDING_TYPE]

def get_water_standards(building_type=None):
    """คืนค่ามาตรฐานน้ำตามประเภทอาคาร (default = ก)"""
    bt = (building_type or DEFAULT_BUILDING_TYPE).strip()
    return WATER_STANDARDS_BY_TYPE.get(bt, WATER_STANDARDS_BY_TYPE[DEFAULT_BUILDING_TYPE])

def find_std_for_param(param_key, building_type=None):
    """หา std dict ของ parameter (รองรับการเลือกตามประเภทอาคาร)"""
    standards = get_water_standards(building_type)
    nk = normalize_key(param_key)
    for k, v in standards.items():
        if normalize_key(k) == nk: return v
    return None

def standard_text(std):
    if not std: return '-'
    t = std.get('type')
    if t == 'max': return f"≤ {std.get('max')}"
    if t == 'range': return f"{std.get('min')} – {std.get('max')}"
    if std.get('min') is not None and std.get('max') is not None: return f"{std.get('min')} – {std.get('max')}"
    if std.get('max') is not None: return f"≤ {std.get('max')}"
    return '-'

def classify_compliance(value, std, tol=0.10):
    if value is None or not std: return 'N/A'
    t = std.get('type')
    if t == 'max':
        mx = std.get('max')
        if mx is None: return 'N/A'
        if value <= mx: return 'Pass'
        if value <= mx * (1 + tol): return 'Near Limit'
        return 'Fail'
    if t == 'range':
        mn, mx = std.get('min'), std.get('max')
        if mn is None or mx is None: return 'N/A'
        if mn <= value <= mx: return 'Pass'
        if mn*(1-tol) <= value <= mx*(1+tol): return 'Near Limit'
        return 'Fail'
    return 'N/A'

def build_parameters(row):
    def get_attr_any(obj, names):
        for n in names:
            if hasattr(obj, n):
                v = getattr(obj, n)
                if v is not None: return v
        for n in names:
            if hasattr(obj, n): return getattr(obj, n)
        return None
    return {
        'pH': safe_float(get_attr_any(row,['pH'])),
        'BOD': safe_float(get_attr_any(row,['BOD'])),
        'COD': safe_float(get_attr_any(row,['COD'])),
        'TSS': safe_float(get_attr_any(row,['TSS'])),
        'TDS': safe_float(get_attr_any(row,['TDS'])),
        'O&G': safe_float(get_attr_any(row,['oil_grease','O_G','OG','OandG','O&G'])),
        'TKN': safe_float(get_attr_any(row,['TKN'])),
        'Sulfide': safe_float(get_attr_any(row,['sulfide','Sulfide'])),
        'TCB': safe_float(get_attr_any(row,['TCB'])),
        'FCB': safe_float(get_attr_any(row,['FCB'])),
        'Cl2': safe_float(get_attr_any(row,['chlorine','Cl2','CL2'])),
    }

def can_fuzzy_eval(params):
    return all(params.get(k) is not None for k in REQUIRED_FUZZY_COLS)

def rule_based_status(param_name, value, building_type=None):
    return classify_compliance(value, find_std_for_param(param_name, building_type), tol=0.10)

def percent_reduction(before, after):
    if before is None or after is None: return None
    try:
        b, a = float(before), float(after)
        if b <= 0: return None
        return round((b - a) / b * 100.0, 1)
    except: return None

def advanced_water_analysis(parameter_results, params):
    analysis = {"top_pollutants":[],"risk_index":"Low","risk_score":0,
                 "compliance_rate":0,"pass_count":0,"near_count":0,"fail_count":0,"total_params":0}
    pollutants, pass_c, near_c, fail_c = [], 0, 0, 0
    for param, result in (parameter_results or {}).items():
        if not isinstance(result, dict): continue
        status = result.get("status")
        # ข้ามพารามิเตอร์ที่ "ไม่กำหนด" ทั้งหมด (regulated=False หรือ risk=None)
        if status == "ไม่กำหนด" or result.get("regulated") is False:
            continue
        risk = result.get("risk_percent")
        # ถ้า risk เป็น None อาจมาจาก fuzzy ไม่สามารถคำนวณได้ — ตีเป็น 0
        if risk is None:
            risk = 0
        try:
            risk = float(risk)
        except (TypeError, ValueError):
            risk = 0.0
        pollutants.append({"name":param,"risk":risk,"status":status})
        if status == "Pass": pass_c += 1
        elif status == "Near Limit": near_c += 1
        elif status == "Fail": fail_c += 1
    pollutants.sort(key=lambda x: x["risk"], reverse=True)
    analysis["top_pollutants"] = pollutants[:3]
    total = pass_c + near_c + fail_c
    analysis.update({"pass_count":pass_c,"near_count":near_c,"fail_count":fail_c,"total_params":total})
    if total > 0: analysis["compliance_rate"] = round(pass_c/total*100,1)
    risk_score = sum(p["risk"] for p in pollutants)/len(pollutants) if pollutants else 0
    analysis["risk_score"] = round(risk_score, 1)
    if risk_score < 25: analysis["risk_index"] = "Low Risk"
    elif risk_score < 50: analysis["risk_index"] = "Moderate Risk"
    elif risk_score < 75: analysis["risk_index"] = "High Risk"
    else: analysis["risk_index"] = "Critical Risk"
    return analysis

def calculate_detailed_statistics(location, pond, date_from, date_to):
    data_list = WaterQualityData.query.filter(
        WaterQualityData.sample_location == location,
        WaterQualityData.pond_name == pond,
        WaterQualityData.sample_date.between(date_from, date_to)
    ).order_by(WaterQualityData.sample_date.asc()).all()
    if not data_list: return None
    param_stats = {}
    parameters = {
        'pH':[d.pH for d in data_list if d.pH is not None],
        'BOD':[d.BOD for d in data_list if d.BOD is not None],
        'COD':[d.COD for d in data_list if d.COD is not None],
        'TSS':[d.TSS for d in data_list if d.TSS is not None],
        'TDS':[d.TDS for d in data_list if d.TDS is not None],
        'O&G':[d.oil_grease for d in data_list if d.oil_grease is not None],
        'TKN':[d.TKN for d in data_list if d.TKN is not None],
        'Sulfide':[d.sulfide for d in data_list if d.sulfide is not None],
        'TCB':[d.TCB for d in data_list if d.TCB is not None],
        'FCB':[d.FCB for d in data_list if d.FCB is not None],
        'Cl2':[d.chlorine for d in data_list if d.chlorine is not None]
    }
    for param_name, values in parameters.items():
        if not values: continue
        standard = find_std_for_param(param_name) or {}
        pass_count = sum(1 for v in values if rule_based_status(param_name, v) == 'Pass')
        near_count = sum(1 for v in values if rule_based_status(param_name, v) == 'Near Limit')
        fail_count = sum(1 for v in values if rule_based_status(param_name, v) == 'Fail')
        total_count = len(values)
        trend = 'stable'
        if len(values) >= 4:
            mid = len(values)//2
            f_avg = sum(values[:mid])/mid
            s_avg = sum(values[mid:])/(len(values)-mid)
            if param_name not in ['pH','Cl2']:
                if s_avg < f_avg*0.9: trend = 'improving'
                elif s_avg > f_avg*1.1: trend = 'worsening'
        param_stats[param_name] = {
            'name':standard.get('name',param_name),'unit':standard.get('unit',''),
            'current':round(values[-1],2),'average':round(sum(values)/len(values),2),
            'min':round(min(values),2),'max':round(max(values),2),
            'standard_min':standard.get('min'),'standard_max':standard.get('max'),
            'pass_rate':round(pass_count/total_count*100,1),'pass_count':pass_count,
            'near_limit_count':near_count,'fail_count':fail_count,
            'total_count':total_count,'trend':trend,
            'status':rule_based_status(param_name,values[-1]),'values':values
        }
    return param_stats

def calculate_all_ponds_overview(location, pond_list, date_from, date_to, calculation_mode='latest'):
    all_ponds_data = []
    for pond_name in pond_list:
        data_list = WaterQualityData.query.filter(
            WaterQualityData.sample_location == location,
            WaterQualityData.pond_name == pond_name,
            WaterQualityData.sample_date.between(date_from, date_to)
        ).order_by(WaterQualityData.sample_date.asc()).all()
        if not data_list: continue
        mapping = {'pH':'pH','BOD':'BOD','COD':'COD','TSS':'TSS','TDS':'TDS',
                   'O&G':'oil_grease','TKN':'TKN','Sulfide':'sulfide',
                   'TCB':'TCB','FCB':'FCB','Cl2':'chlorine'}
        if calculation_mode == 'latest':
            latest_data = data_list[-1]
            parameters = build_parameters(latest_data)
            sample_date = latest_data.sample_date
            calculation_info = f"ข้อมูลล่าสุด ({latest_data.sample_date.strftime('%d/%m/%Y')})"
        elif calculation_mode == 'average':
            parameters = {}
            for k, attr in mapping.items():
                vals = [safe_float(getattr(d,attr)) for d in data_list if safe_float(getattr(d,attr)) is not None]
                parameters[k] = sum(vals)/len(vals) if vals else None
            sample_date = data_list[-1].sample_date
            calculation_info = f"ค่าเฉลี่ย {len(data_list)} วัน"
        else:
            parameters = {}
            now = datetime.now()
            for k, attr in mapping.items():
                ws = tw = 0.0
                for d in data_list:
                    v = safe_float(getattr(d,attr))
                    if v is None: continue
                    w = 1.0/((now - d.sample_date).days + 1.0)
                    ws += v*w; tw += w
                parameters[k] = ws/tw if tw > 0 else None
            sample_date = data_list[-1].sample_date
            calculation_info = f"ค่าเฉลี่ยถ่วงน้ำหนัก {len(data_list)} วัน"

        if can_fuzzy_eval(parameters):
            # ใช้ building_type ของข้อมูลล่าสุดในกลุ่มนี้
            _bt = getattr(data_list[-1], 'building_type', None) or DEFAULT_BUILDING_TYPE
            fuzzy_eval = get_fuzzy_system(_bt).evaluate_overall(parameters)
            score = fuzzy_eval.get('overall_score')
            overall_status = score_to_overall_status(score)
        else:
            fuzzy_eval = {'overall_score':None,'overall_level':'No Data',
                          'overall_message':'ข้อมูลไม่ครบสำหรับประเมิน Fuzzy',
                          'parameter_results':{},'fuzzy_membership':{}}
            score = None; overall_status = 'No Data'

        critical_params, warning_params = [], []
        for param, result in (fuzzy_eval.get('parameter_results') or {}).items():
            if isinstance(result, dict):
                st = result.get('status')
                if st == 'Fail': critical_params.append(param)
                elif st == 'Near Limit': warning_params.append(param)

        all_ponds_data.append({
            'pond_name':pond_name,'status':overall_status,'fuzzy_score':score,
            'fuzzy_message':fuzzy_eval.get('overall_message',''),'sample_date':sample_date,
            'critical_count':0,'warning_count':0,'pass_count':0,'total_parameters':0,
            'critical_params':critical_params,'warning_params':warning_params,
            'parameter_results':fuzzy_eval.get('parameter_results',{}),
            'calculation_mode':calculation_mode,'calculation_info':calculation_info,
        })
    return all_ponds_data

def is_final_pond(name):
    n = (name or "").lower()
    return any(k in n for k in ["หลังบำบัด","treated","final","effluent","บ่อพักน้ำทิ้ง","treated water tank"])

def _day_range(dt):
    return (dt.replace(hour=0,minute=0,second=0,microsecond=0),
            dt.replace(hour=23,minute=59,second=59,microsecond=0))

def pick_latest_by_keywords_in_day(loc, sample_type, keywords, day_start, day_end):
    q = WaterQualityData.query.filter(
        WaterQualityData.sample_location == loc,
        WaterQualityData.sample_type == sample_type,
        WaterQualityData.sample_date.between(day_start, day_end)
    )
    cond = None
    for kw in keywords:
        c = WaterQualityData.pond_name.ilike(f"%{kw}%")
        cond = c if cond is None else (cond | c)
    if cond is not None: q = q.filter(cond)
    return q.order_by(WaterQualityData.sample_date.desc()).first()

# =========================
# AUTO FUZZY SYNC
# =========================

class SystemConfig(db.Model):
    __tablename__ = 'system_config'
    id    = db.Column(db.Integer, primary_key=True)
    key   = db.Column(db.String(64), unique=True, nullable=False)
    value = db.Column(db.String(255), nullable=False)


def get_fuzzy_version() -> str:
    config_snapshot = {
        'fuzzy_version': FUZZY_VERSION,
        'standards': {
            k: {'min': v.get('min'), 'max': v.get('max'), 'type': v.get('type')}
            for k, v in WATER_STANDARDS.items()
        },
    }
    raw = json.dumps(config_snapshot, sort_keys=True, ensure_ascii=False)
    return hashlib.md5(raw.encode()).hexdigest()[:10]


def _re_evaluate_all_rows(location_filter: str = '') -> dict:
    query = WaterQualityData.query.filter(WaterQualityData.sample_type == 'หลังบำบัด')
    if location_filter:
        query = query.filter(WaterQualityData.sample_location == location_filter)
    rows = query.all()
    updated = skipped = errors = 0
    for data in rows:
        try:
            params = build_parameters(data)
            if not can_fuzzy_eval(params):
                skipped += 1; continue
            fe     = get_fuzzy_system(getattr(data, 'building_type', None) or DEFAULT_BUILDING_TYPE).evaluate_overall(params)
            score  = fe.get('overall_score')
            status = score_to_overall_status(score)
            pr     = fe.get('parameter_results', {}) or {}
            result = EvaluationResult.query.filter_by(water_data_id=data.id).first()
            if result:
                result.overall_score  = score
                result.overall_status = status
                result.overall_message = fe.get('overall_message', '')
            else:
                result = EvaluationResult(
                    water_data_id=data.id, overall_score=score,
                    overall_status=status, overall_message=fe.get('overall_message', ''),
                    pass_count=0, near_limit_count=0, fail_count=0, total_parameters=0,
                )
                db.session.add(result)
            for attr, param_key in [
                ('pH_result','pH'),('BOD_result','BOD'),('COD_result','COD'),
                ('TSS_result','TSS'),('TDS_result','TDS'),('oil_grease_result','O&G'),
                ('TKN_result','TKN'),('sulfide_result','Sulfide'),
                ('TCB_result','TCB'),('FCB_result','FCB'),('chlorine_result','Cl2'),
            ]:
                if hasattr(result, attr):
                    setattr(result, attr, pr.get(param_key))
            updated += 1
        except Exception as e:
            errors += 1
            app.logger.warning(f'[re-evaluate] id={data.id}: {e}')
    db.session.commit()
    return {'updated': updated, 'skipped': skipped, 'errors': errors}


def check_and_sync_fuzzy():
    try:
        current_ver = get_fuzzy_version()
        saved = SystemConfig.query.filter_by(key='fuzzy_version').first()
        saved_ver = saved.value if saved else None
        if saved_ver == current_ver:
            app.logger.info(f'[Fuzzy] config ไม่เปลี่ยน ({current_ver}) — ข้าม re-evaluate')
            return
        app.logger.info(f'[Fuzzy] config เปลี่ยน: {saved_ver!r} → {current_ver!r} — เริ่ม re-evaluate...')
        result = _re_evaluate_all_rows()
        app.logger.info(f'[Fuzzy] เสร็จ · อัปเดต {result["updated"]} · ข้าม {result["skipped"]} · ผิดพลาด {result["errors"]}')
        if saved: saved.value = current_ver
        else: db.session.add(SystemConfig(key='fuzzy_version', value=current_ver))
        db.session.commit()
    except Exception as e:
        app.logger.error(f'[Fuzzy] check_and_sync_fuzzy failed: {e}')


# =========================
# LOGIN
# =========================
@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

@app.route('/')
def index():
    return redirect(url_for('public_dashboard'))

@app.route('/login', methods=['GET','POST'])
def login():
    if current_user.is_authenticated: return redirect(url_for('dashboard'))
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user)
            flash(f'ยินดีต้อนรับ {user.full_name or user.username}!', 'success')
            return redirect(request.args.get('next') or url_for('dashboard'))
        flash('ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง', 'error')
    return render_template('login.html')

@app.route('/logout')
@login_required
def logout():
    logout_user()
    flash('ออกจากระบบเรียบร้อย', 'success')
    return redirect(url_for('login'))

# =========================
# PUBLIC DASHBOARD
# =========================
@app.route('/public/dashboard')
def public_dashboard():
    latest_subq = db.session.query(
        WaterQualityData.sample_location.label('loc'),
        WaterQualityData.pond_name.label('pond'),
        func.max(WaterQualityData.sample_date).label('max_date')
    ).filter(WaterQualityData.sample_type == 'หลังบำบัด').group_by(
        WaterQualityData.sample_location, WaterQualityData.pond_name
    ).subquery()

    latest_rows = db.session.query(WaterQualityData).join(
        latest_subq,
        (WaterQualityData.sample_location == latest_subq.c.loc) &
        (WaterQualityData.pond_name == latest_subq.c.pond) &
        (WaterQualityData.sample_date == latest_subq.c.max_date)
    ).all()

    # ใช้ครั้งละมาตรฐานตามประเภทอาคารของ row นั้น (ไม่ใช่ค่า hardcoded เดิม)
    all_ponds = []
    crisis_params = []

    for data in latest_rows:
        params = build_parameters(data)
        bt = getattr(data, 'building_type', None) or DEFAULT_BUILDING_TYPE
        if can_fuzzy_eval(params):
            fe     = get_fuzzy_system(bt).evaluate_overall(params)
            score  = fe.get('overall_score')
            status = score_to_overall_status(score)
        else:
            score  = None
            status = 'No Data'

        all_ponds.append({
            'location':    data.sample_location,
            'pond_name':   data.pond_name,
            'sample_date': data.sample_date,
            'sample_type': data.sample_type,
            'status':      status,
            'fuzzy_score': score,
            'pH':  data.pH, 'BOD': data.BOD, 'COD': data.COD,
            'TSS': data.TSS, 'TDS': data.TDS, 'FCB': data.FCB,
        })

        # crisis params: ใช้มาตรฐานของอาคารตัวเอง
        row_stds = BUILDING_STANDARDS.get(bt, BUILDING_STANDARDS[DEFAULT_BUILDING_TYPE])
        for param in ['BOD', 'COD', 'TSS', 'TDS', 'FCB']:
            std_dict = row_stds.get(param)
            if std_dict is None:
                continue  # ไม่กำหนดสำหรับอาคารนี้
            std_max = std_dict.get('max')
            if std_max is None:
                continue
            val = getattr(data, param, None)
            if val is not None and val > std_max * 2:
                crisis_params.append({'pond': data.pond_name, 'param': param,
                                      'value': round(val, 1), 'times': round(val / std_max, 1)})

    total        = len(all_ponds)
    valid_scores = [p['fuzzy_score'] for p in all_ponds if p.get('fuzzy_score') is not None]
    overall_stats = {
        'total':      total,
        'pass':       sum(1 for p in all_ponds if p.get('status') == 'Pass'),
        'near_limit': sum(1 for p in all_ponds if p.get('status') == 'Near Limit'),
        'fail':       sum(1 for p in all_ponds if p.get('status') == 'Fail'),
        'avg_score':  round(sum(valid_scores) / len(valid_scores), 1) if valid_scores else 0,
        'pass_rate':  round(sum(1 for p in all_ponds if p.get('status') == 'Pass') / total * 100, 1) if total > 0 else 0,
    }

    seven_days_ago = datetime.now() - timedelta(days=7)
    trend_data = WaterQualityData.query.filter(
        WaterQualityData.sample_type == 'หลังบำบัด',
        WaterQualityData.sample_date >= seven_days_ago
    ).order_by(WaterQualityData.sample_date.asc()).all()

    trend_by_date = {}
    for data in trend_data:
        d = data.sample_date.strftime('%d/%m')
        trend_by_date.setdefault(d, [])
        params = build_parameters(data)
        if can_fuzzy_eval(params):
            score = get_fuzzy_system(getattr(data, 'building_type', None) or DEFAULT_BUILDING_TYPE).evaluate_overall(params).get('overall_score')
            if score is not None: trend_by_date[d].append(score)

    trend_labels, trend_scores = [], []
    for d in sorted(trend_by_date.keys()):
        scores = [s for s in trend_by_date[d] if s is not None]
        if scores:
            trend_labels.append(d)
            trend_scores.append(round(sum(scores) / len(scores), 1))

    thirty_ago = datetime.now() - timedelta(days=30)
    recent_rows = WaterQualityData.query.filter(
        WaterQualityData.sample_type == 'หลังบำบัด',
        WaterQualityData.sample_date >= thirty_ago,
    ).order_by(WaterQualityData.sample_date.desc()).limit(1000).all()

    all_dates  = set(r.sample_date.date() for r in recent_rows)
    pass_dates = set()
    for r in recent_rows:
        params = build_parameters(r)
        if can_fuzzy_eval(params):
            sc = get_fuzzy_system(getattr(r, 'building_type', None) or DEFAULT_BUILDING_TYPE).evaluate_overall(params).get('overall_score')
            if sc is not None and score_to_overall_status(sc) == 'Pass':
                pass_dates.add(r.sample_date.date())

    monthly_total_days = len(all_dates)
    monthly_pass_days  = len(pass_dates)
    monthly_pass_rate  = round(monthly_pass_days / monthly_total_days * 100) if monthly_total_days else 0

    from collections import Counter
    param_exceed = Counter()
    param_total  = Counter()
    # ใช้มาตรฐานตามประเภทอาคารของแต่ละ row
    for r in recent_rows:
        bt = getattr(r, 'building_type', None) or DEFAULT_BUILDING_TYPE
        row_stds = BUILDING_STANDARDS.get(bt, BUILDING_STANDARDS[DEFAULT_BUILDING_TYPE])
        for param in ['BOD', 'COD', 'TSS', 'TDS', 'FCB']:
            std_dict = row_stds.get(param)
            if std_dict is None:
                continue  # ไม่กำหนดสำหรับอาคารนี้
            std_max = std_dict.get('max')
            if std_max is None:
                continue
            val = getattr(r, param, None)
            if val is not None:
                param_total[param] += 1
                if val > std_max:
                    param_exceed[param] += 1

    top_problem_params = sorted([
        {'param': p, 'rate': round(param_exceed[p] / param_total[p] * 100) if param_total[p] else 0,
         'days': param_exceed[p]}
        for p in ['BOD', 'COD', 'TSS', 'TDS', 'FCB'] if param_total[p] > 0
    ], key=lambda x: x['rate'], reverse=True)

    return render_template(
        'public_dashboard.html',
        all_ponds=all_ponds, overall_stats=overall_stats,
        trend_labels=trend_labels, trend_scores=trend_scores,
        water_standards=to_object(WATER_STANDARDS), last_update=datetime.now(),
        crisis_params=crisis_params,
        monthly_pass_rate=monthly_pass_rate, monthly_pass_days=monthly_pass_days,
        monthly_total_days=monthly_total_days, top_problem_params=top_problem_params,
    )

# =========================
# MAIN DASHBOARD
# =========================
@app.route('/dashboard')
@login_required
def dashboard():
    locations = db.session.query(WaterQualityData.sample_location).distinct().filter(
        WaterQualityData.sample_location.isnot(None), WaterQualityData.sample_location != ''
    ).all()
    location_list = [loc[0] for loc in locations if loc[0]]

    latest_loc = db.session.query(WaterQualityData.sample_location).filter(
        WaterQualityData.sample_type == 'หลังบำบัด',
        WaterQualityData.sample_location.isnot(None), WaterQualityData.sample_location != ''
    ).order_by(WaterQualityData.sample_date.desc()).first()

    default_location = latest_loc[0] if latest_loc else (location_list[0] if location_list else None)
    selected_location = request.args.get('location', default_location)

    ponds = db.session.query(WaterQualityData.pond_name).distinct().filter(
        WaterQualityData.sample_location == selected_location
    ).all()
    pond_list = [p[0] for p in ponds if p[0]]
    selected_pond = request.args.get('pond', 'ทั้งหมด') or 'ทั้งหมด'

    mode = (request.args.get('mode') or 'latest').lower()
    month_arg = request.args.get('month')
    date_from_arg = request.args.get('date_from')
    date_to_arg = request.args.get('date_to')

    recent_days = None
    if mode == 'range' and date_from_arg and date_to_arg:
        try:
            _df = datetime.strptime(date_from_arg, '%Y-%m-%d')
            _dt = datetime.strptime(date_to_arg, '%Y-%m-%d')
            _today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
            if _dt.date() == _today.date():
                _diff = (_dt - _df).days
                if _diff in (7, 15, 30): recent_days = _diff
        except Exception:
            pass
    selected_month = month_arg or datetime.now().strftime('%Y-%m')

    if date_from_arg and date_to_arg and mode == 'latest':
        mode = 'range'

    def month_start_end(yyyy_mm):
        y, m = map(int, yyyy_mm.split('-'))
        start = datetime(y, m, 1, 0, 0, 0)
        next_m = datetime(y+1, 1, 1) if m == 12 else datetime(y, m+1, 1)
        return start, next_m - timedelta(seconds=1)

    date_mode = 'latest'
    date_from_str = date_to_str = display_date_text = ''

    if mode == 'range' and date_from_arg and date_to_arg:
        date_from = datetime.strptime(date_from_arg, '%Y-%m-%d')
        date_to = datetime.strptime(date_to_arg, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
        date_mode = 'range'
        date_from_str, date_to_str = date_from_arg, date_to_arg
        display_date_text = f"แสดงผลจากวันที่ {date_from_arg} ถึง {date_to_arg}"
    elif mode == 'month':
        try:
            date_from, date_to = month_start_end(selected_month)
            date_mode = 'month'
            date_from_str = date_from.strftime('%Y-%m-%d')
            date_to_str = date_to.strftime('%Y-%m-%d')
            display_date_text = f"แสดงผลแบบเดือน: {selected_month}"
        except:
            mode = 'latest'

    if mode in ('latest', 'range') and not date_from_str:
        latest_record = WaterQualityData.query.filter(
            WaterQualityData.sample_location == selected_location,
            WaterQualityData.sample_type == 'หลังบำบัด'
        ).order_by(WaterQualityData.sample_date.desc()).first()
        if latest_record:
            date_from = latest_record.sample_date.replace(hour=0, minute=0, second=0, microsecond=0)
            date_to = latest_record.sample_date.replace(hour=23, minute=59, second=59, microsecond=0)
            date_from_str = latest_record.sample_date.strftime('%Y-%m-%d')
            date_to_str = latest_record.sample_date.strftime('%Y-%m-%d')
            date_mode = 'latest'
            display_date_text = f"แสดงผลจากข้อมูลล่าสุด ณ วันที่ {latest_record.sample_date.strftime('%d/%m/%Y %H:%M')}"
        else:
            date_from = date_to = datetime.now()
            display_date_text = "ไม่พบข้อมูลล่าสุด"

    _avail_q = db.session.query(func.date(WaterQualityData.sample_date)).filter(
        WaterQualityData.sample_location == selected_location,
        WaterQualityData.sample_date.between(date_from, date_to)
    )
    if selected_pond != 'ทั้งหมด':
        _avail_q = _avail_q.filter(WaterQualityData.pond_name == selected_pond)
    available_dates_q = _avail_q.distinct().order_by(func.date(WaterQualityData.sample_date).desc()).all()
    available_dates = [str(d[0]) for d in available_dates_q]

    snapshot_date = request.args.get('snapshot_date')
    if not snapshot_date or snapshot_date not in available_dates:
        snapshot_date = available_dates[0] if available_dates else None

    if snapshot_date:
        snap_start = datetime.strptime(snapshot_date, '%Y-%m-%d').replace(hour=0, minute=0, second=0, microsecond=0)
        snap_end   = datetime.strptime(snapshot_date, '%Y-%m-%d').replace(hour=23, minute=59, second=59, microsecond=0)
    else:
        snap_start = date_from
        snap_end   = date_to

    if selected_pond == 'ทั้งหมด':
        calculation_mode = request.args.get('calc_mode', 'latest')
        all_ponds_overview = calculate_all_ponds_overview(
            selected_location, pond_list, date_from, date_to, calculation_mode)
        all_ponds_overview.sort(key=lambda p: (
            0 if is_final_pond(p.get('pond_name')) else 1,
            -(p.get('fuzzy_score') or -1),
            -(p.get('sample_date').timestamp() if p.get('sample_date') else 0)
        ))

        overview_stats = {
            'total': len(all_ponds_overview),
            'pass': sum(1 for p in all_ponds_overview if p.get('status') == 'Pass'),
            'near_limit': sum(1 for p in all_ponds_overview if p.get('status') == 'Near Limit'),
            'fail': sum(1 for p in all_ponds_overview if p.get('status') == 'Fail')
        }

        stage_compare, reduction, match_date_text = [], {}, None
        final_item = next((p for p in all_ponds_overview if is_final_pond(p.get('pond_name'))), None)

        if snapshot_date:
            day_start, day_end = snap_start, snap_end
            match_date_text = snap_start.strftime('%d/%m/%Y')
        elif final_item and final_item.get('sample_date'):
            matched_day_dt = final_item.get('sample_date')
            day_start, day_end = _day_range(matched_day_dt)
            match_date_text = matched_day_dt.strftime('%d/%m/%Y')
        else:
            day_start = day_end = None

        if day_start:
            eq_row = pick_latest_by_keywords_in_day(selected_location, 'ก่อนบำบัด',
                ['บ่อรวบรวมนํ้าเสีย','บ่อรวบรวมน้ำเสีย','Equalization','Collection','รวบรวม'], day_start, day_end)
            aer_row = pick_latest_by_keywords_in_day(selected_location, 'กำลังบำบัด',
                ['บ่อเติมอากาศ','Aeration','เติมอากาศ'], day_start, day_end)
            sed_row = pick_latest_by_keywords_in_day(selected_location, 'กำลังบำบัด',
                ['บ่อตกตะกอน','Sedimentation','ตกตะกอน'], day_start, day_end)
            final_row = WaterQualityData.query.filter(
                WaterQualityData.sample_location == selected_location,
                WaterQualityData.sample_type == 'หลังบำบัด',
                WaterQualityData.sample_date.between(day_start, day_end)
            ).order_by(WaterQualityData.sample_date.desc()).first()

            def stage_row_dict(title, row):
                if not row: return None
                p = build_parameters(row)
                fr = (get_fuzzy_system(getattr(row,'building_type',None) or DEFAULT_BUILDING_TYPE).evaluate_overall(p)
                      if can_fuzzy_eval(p) else {})
                score = fr.get('overall_score')
                return {'stage': title, 'pond_name': getattr(row,'pond_name','-'),
                        'sample_type': getattr(row,'sample_type','-'),
                        'BOD': p.get('BOD'), 'COD': p.get('COD'), 'TSS': p.get('TSS'),
                        'FCB': p.get('FCB'), 'TCB': p.get('TCB'),
                        'score': score, 'status': score_to_overall_status(score)}

            for title, row in [('ก่อนบำบัด (บ่อรวบรวม)', eq_row),
                                ('กำลังบำบัด (เติมอากาศ)', aer_row),
                                ('กำลังบำบัด (ตกตะกอน)', sed_row),
                                ('หลังบำบัด (บ่อพักน้ำทิ้ง)', final_row)]:
                sr = stage_row_dict(title, row)
                if sr: stage_compare.append(sr)

            if eq_row and final_row:
                b, a = build_parameters(eq_row), build_parameters(final_row)
                reduction = {
                    'BOD': percent_reduction(b.get('BOD'), a.get('BOD')),
                    'COD': percent_reduction(b.get('COD'), a.get('COD')),
                    'TSS': percent_reduction(b.get('TSS'), a.get('TSS')),
                    'FCB': percent_reduction(b.get('FCB'), a.get('FCB')),
                    'TCB': percent_reduction(b.get('TCB'), a.get('TCB')),
                }

        all_stage_data = {}
        for _d in available_dates[:60]:
            _ds = datetime.strptime(_d, '%Y-%m-%d').replace(hour=0, minute=0, second=0)
            _de = datetime.strptime(_d, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
            _eq  = pick_latest_by_keywords_in_day(selected_location, 'ก่อนบำบัด',
                ['บ่อรวบรวมนํ้าเสีย','บ่อรวบรวมน้ำเสีย','Equalization','Collection','รวบรวม'], _ds, _de)
            _aer = pick_latest_by_keywords_in_day(selected_location, 'กำลังบำบัด',
                ['บ่อเติมอากาศ','Aeration','เติมอากาศ'], _ds, _de)
            _sed = pick_latest_by_keywords_in_day(selected_location, 'กำลังบำบัด',
                ['บ่อตกตะกอน','Sedimentation','ตกตะกอน'], _ds, _de)
            _fin = WaterQualityData.query.filter(
                WaterQualityData.sample_location == selected_location,
                WaterQualityData.sample_type == 'หลังบำบัด',
                WaterQualityData.sample_date.between(_ds, _de)
            ).order_by(WaterQualityData.sample_date.desc()).first()

            _day_stages = []
            for _title, _row in [('ก่อนบำบัด (บ่อรวบรวม)',_eq),('กำลังบำบัด (เติมอากาศ)',_aer),
                                  ('กำลังบำบัด (ตกตะกอน)',_sed),('หลังบำบัด (บ่อพักน้ำทิ้ง)',_fin)]:
                if not _row: continue
                _p = build_parameters(_row)
                _fr = (get_fuzzy_system(getattr(_row,'building_type',None) or DEFAULT_BUILDING_TYPE).evaluate_overall(_p)
                       if can_fuzzy_eval(_p) else {})
                _sc = _fr.get('overall_score')
                _day_stages.append({'stage':_title,'pond_name':getattr(_row,'pond_name','-'),
                    'sample_type':getattr(_row,'sample_type','-'),
                    'BOD':_p.get('BOD'),'COD':_p.get('COD'),'TSS':_p.get('TSS'),'FCB':_p.get('FCB'),'score':_sc})
            _red = {}
            if _eq and _fin:
                _b, _a = build_parameters(_eq), build_parameters(_fin)
                _red = {k: percent_reduction(_b.get(k), _a.get(k)) for k in ['BOD','COD','TSS','FCB','TCB']}
            all_stage_data[_d] = {'stages': _day_stages, 'reduction': _red}

        trend_labels, trend_bod, trend_cod, trend_tss, trend_ph, trend_coliform = [], [], [], [], [], []
        _show_trend = (date_from_str != date_to_str)
        if _show_trend:
            _latest_in_filter = WaterQualityData.query.filter(
                WaterQualityData.sample_location == selected_location,
                WaterQualityData.sample_date.between(date_from, date_to)
            ).order_by(WaterQualityData.sample_date.desc()).first()
            if _latest_in_filter:
                _trend_anchor = _latest_in_filter.sample_date
                seven_days_ago = _trend_anchor - timedelta(days=7)
                _trend_raw = WaterQualityData.query.filter(
                    WaterQualityData.sample_location == selected_location,
                    WaterQualityData.sample_date.between(seven_days_ago, _trend_anchor)
                ).order_by(WaterQualityData.sample_date.asc()).all()
                _after_raw = [r for r in _trend_raw if r.sample_type == 'หลังบำบัด']
                _use_raw = _after_raw if _after_raw else _trend_raw
                _trend_by_date = defaultdict(list)
                for _r in _use_raw:
                    _trend_by_date[_r.sample_date.strftime('%d/%m/%Y')].append(_r)
                for _d in sorted(_trend_by_date.keys()):
                    _rows = _trend_by_date[_d]
                    def _avg(attr, rows=_rows):
                        vals = [safe_float(getattr(r,attr)) for r in rows if safe_float(getattr(r,attr)) is not None]
                        return round(sum(vals)/len(vals),2) if vals else None
                    trend_labels.append(_d[:-5])
                    trend_bod.append(_avg('BOD')); trend_cod.append(_avg('COD'))
                    trend_tss.append(_avg('TSS')); trend_ph.append(_avg('pH'))
                    trend_coliform.append(_avg('FCB'))
                if len(trend_labels) < 2:
                    trend_labels, trend_bod, trend_cod, trend_tss, trend_ph, trend_coliform = [], [], [], [], [], []

        thirty_days_ago = datetime.now() - timedelta(days=30)
        _events_raw = WaterQualityData.query.filter(
            WaterQualityData.sample_location == selected_location,
            WaterQualityData.sample_date >= thirty_days_ago,
            WaterQualityData.sample_type == 'หลังบำบัด'  # ✅ เฉพาะน้ำออก
        ).order_by(WaterQualityData.sample_date.desc()).limit(300).all()

        _PARAM_MAP = [('BOD','BOD','mg/L'),('COD','COD','mg/L'),('TSS','TSS','mg/L'),
                      ('pH','pH',''),('TCB','TCB','MPN/100mL'),('FCB','FCB','MPN/100mL'),
                      ('TDS','TDS','mg/L'),('TKN','TKN','mg/L')]
        
        # ✅ จัดกลุ่มตาม (วันที่+เวลา, บ่อ) — แถวเดียวกัน = วันเดียวบ่อเดียว
        _ev_map = {}
        for _r in _events_raw:
            _key = (_r.sample_date.strftime('%d/%m/%Y %H:%M'), _r.pond_name or '-')
            if _key not in _ev_map:
                if len(_ev_map) >= 20: continue
                _ev_map[_key] = {
                    'date': _r.sample_date.strftime('%d/%m/%Y %H:%M'),
                    'pond': _r.pond_name or '-',
                    'date_for_url': _r.sample_date.strftime('%Y-%m-%d'),
                    'pond_url': _r.pond_name or '',
                    'params': [],
                    'has_fail': False,
                }
            _entry = _ev_map[_key]
            for _param, _attr, _unit in _PARAM_MAP:
                if any(p['param'] == _param for p in _entry['params']): continue
                _val = safe_float(getattr(_r, _attr, None))
                if _val is None: continue
                _std = find_std_for_param(_param)
                _comp = classify_compliance(_val, _std, tol=0.10)
                if _comp in ('Fail', 'Near Limit'):
                    _entry['params'].append({
                        'param': _param,
                        'status': 'Fail' if _comp == 'Fail' else 'Near',
                    })
                    if _comp == 'Fail': _entry['has_fail'] = True
        critical_events = [v for v in _ev_map.values() if v['params']]

        _HM = [('pH','pH'),('BOD','BOD'),('COD','COD'),('TSS','TSS'),('TDS','TDS'),('FCB','FCB'),('TCB','TCB')]
        heatmap_rows = []
        for _pn in pond_list:
            _pr = WaterQualityData.query.filter(
                WaterQualityData.sample_location == selected_location,
                WaterQualityData.pond_name == _pn,
                WaterQualityData.sample_date.between(date_from, date_to)
            ).all()
            if not _pr: continue
            _cell = {}; _tp = _ta = 0
            for _pk, _attr in _HM:
                _vals = [safe_float(getattr(r, _attr)) for r in _pr if safe_float(getattr(r, _attr)) is not None]
                if not _vals: _cell[_pk.lower()] = None; continue
                _p = sum(1 for v in _vals if classify_compliance(v, find_std_for_param(_pk), tol=0.10) == 'Pass')
                _cell[_pk.lower()] = round(_p/len(_vals)*100)
                _tp += _p; _ta += len(_vals)
            heatmap_rows.append({'name':_pn,
                'ph':_cell.get('ph'),'bod':_cell.get('bod'),'cod':_cell.get('cod'),
                'tss':_cell.get('tss'),'tds':_cell.get('tds'),'fcb':_cell.get('fcb'),'tcb':_cell.get('tcb'),
                'avg': round(_tp/_ta*100) if _ta else None})

        _fscores = [p.get('fuzzy_score') for p in all_ponds_overview if p.get('fuzzy_score') is not None]
        _eff_vals = [v for v in [reduction.get('BOD'), reduction.get('COD'), reduction.get('TSS')] if v is not None]
        kpi = {
            'overall_score': round(sum(_fscores)/len(_fscores),1) if _fscores else None,
            'score_trend': '',
            'compliance_rate': round(overview_stats['pass']/max(overview_stats['total'],1)*100),
            'compliance_trend': '',
            'treatment_eff': round(sum(_eff_vals)/len(_eff_vals),1) if _eff_vals else None,
            'eff_trend': '',
            'critical_30d': sum(1 for e in critical_events if e.get('has_fail')),
            'critical_trend': '',
        }

        pond_compare = []
        for _ps in all_ponds_overview:
            _pn = _ps.get('pond_name')
            _rs = WaterQualityData.query.filter(
                WaterQualityData.sample_location == selected_location,
                WaterQualityData.pond_name == _pn,
                WaterQualityData.sample_date.between(date_from, date_to)
            ).all()
            if not _rs: continue
            def _mean(attr, rows=_rs):
                vals = [safe_float(getattr(r,attr)) for r in rows if safe_float(getattr(r,attr)) is not None]
                return round(sum(vals)/len(vals),1) if vals else None
            pond_compare.append({'name':_pn,'bod':_mean('BOD'),'cod':_mean('COD'),
                'tss':_mean('TSS'),'ph':_mean('pH'),
                'score':_ps.get('fuzzy_score'),'status':_ps.get('status','No Data')})

        _bod_vals = [r['bod'] for r in pond_compare if r['bod'] is not None]
        _cod_vals = [r['cod'] for r in pond_compare if r['cod'] is not None]
        _tss_vals = [r['tss'] for r in pond_compare if r['tss'] is not None]
        bod_max_ref = (max(_bod_vals)*1.2) if _bod_vals else 200
        cod_max_ref = (max(_cod_vals)*1.2) if _cod_vals else 300
        tss_max_ref = (max(_tss_vals)*1.2) if _tss_vals else 200

        _all_stat = WaterQualityData.query.filter(
            WaterQualityData.sample_location == selected_location,
            WaterQualityData.sample_date.between(date_from, date_to)
        ).all()
        def _stat(attr):
            vals = [safe_float(getattr(r,attr)) for r in _all_stat if safe_float(getattr(r,attr)) is not None]
            if not vals: return None, None, None
            return round(sum(vals)/len(vals),1), round(max(vals),1), round(min(vals),1)
        bod_mean, bod_max_s, bod_min = _stat('BOD')
        cod_mean, cod_max_s, cod_min = _stat('COD')
        tss_mean, tss_max_s, tss_min = _stat('TSS')

        return render_template('staff_home.html',
            location_list=location_list, pond_list=pond_list,
            selected_location=selected_location, selected_pond=selected_pond,
            all_ponds_overview=all_ponds_overview, overview_stats=overview_stats,
            calculation_mode=calculation_mode, display_date_text=display_date_text,
            date_mode=date_mode, date_from=date_from_str, date_to=date_to_str,
            selected_month=selected_month, stage_compare=stage_compare,
            reduction=reduction, match_date_text=match_date_text,
            water_standards=to_object(WATER_STANDARDS),
            trend_labels=trend_labels, trend_bod=trend_bod, trend_cod=trend_cod,
            trend_tss=trend_tss, trend_ph=trend_ph, trend_coliform=trend_coliform,
            critical_events=critical_events, heatmap_rows=heatmap_rows,
            kpi=kpi, pond_compare=pond_compare,
            bod_max_ref=bod_max_ref, cod_max_ref=cod_max_ref, tss_max_ref=tss_max_ref,
            bod_mean=bod_mean, bod_max=bod_max_s, bod_min=bod_min,
            cod_mean=cod_mean, cod_max=cod_max_s, cod_min=cod_min,
            tss_mean=tss_mean, tss_max=tss_max_s, tss_min=tss_min,
            available_dates=available_dates, snapshot_date=snapshot_date,
            all_stage_data=all_stage_data, recent_days=recent_days,
        )

    # ── DETAIL MODE ──
    after_q = WaterQualityData.query.filter(
        WaterQualityData.sample_location == selected_location,
        WaterQualityData.pond_name == selected_pond,
        WaterQualityData.sample_date.between(snap_start, snap_end)
    ).order_by(WaterQualityData.sample_date.desc())
    after_count = after_q.count()
    after_data  = after_q.first()

    if after_data is None:
        latest_any = WaterQualityData.query.filter(
            WaterQualityData.sample_location == selected_location,
            WaterQualityData.pond_name == selected_pond,
            WaterQualityData.sample_date.between(date_from, date_to)
        ).order_by(WaterQualityData.sample_date.desc()).first()
        if latest_any:
            display_date_text = (f"⚠️ ไม่มีข้อมูลวันที่ {snapshot_date} — "
                f"แสดงล่าสุดแทน ({latest_any.sample_date.strftime('%d/%m/%Y %H:%M')})")
            after_data = latest_any; after_count = 1

    stage_latest = {'equalization':None,'aeration':None,'sedimentation':None,'treated':after_data}
    stage_compare, reduction = [], {}

    def pick_latest_by_keyword(sample_type, keywords, ds, de):
        q = WaterQualityData.query.filter(
            WaterQualityData.sample_location == selected_location,
            WaterQualityData.sample_type == sample_type,
            WaterQualityData.sample_date.between(ds, de)
        )
        cond = None
        for kw in keywords:
            c = WaterQualityData.pond_name.ilike(f"%{kw}%")
            cond = c if cond is None else (cond | c)
        if cond is not None: q = q.filter(cond)
        return q.order_by(WaterQualityData.sample_date.desc()).first()

    if after_data is not None and after_data.sample_date is not None:
        day_start, day_end = snap_start, snap_end
        stage_latest['equalization'] = pick_latest_by_keyword('ก่อนบำบัด',
            ['บ่อรวบรวมนํ้าเสีย','บ่อรวบรวมน้ำเสีย','Equalization','Collection','รวบรวม'], day_start, day_end)
        stage_latest['aeration'] = pick_latest_by_keyword('กำลังบำบัด',
            ['บ่อเติมอากาศ','Aeration','เติมอากาศ'], day_start, day_end)
        stage_latest['sedimentation'] = pick_latest_by_keyword('กำลังบำบัด',
            ['บ่อตกตะกอน','Sedimentation','ตกตะกอน'], day_start, day_end)

    pond_data = {'before':{'latest':stage_latest['equalization']},
                 'inprocess':{'latest':stage_latest['aeration'] or stage_latest['sedimentation']},
                 'after':{'latest':stage_latest['treated']},'stages':stage_latest}

    def stage_row_d(stage_name, sample_type, row):
        if row is None: return None
        p = build_parameters(row)
        score = (get_fuzzy_system(getattr(row,'building_type',None) or DEFAULT_BUILDING_TYPE).evaluate_overall(p).get('overall_score')
                 if can_fuzzy_eval(p) else None)
        return {'stage':stage_name,'sample_type':sample_type,'pond_name':getattr(row,'pond_name','-'),
                'BOD':p.get('BOD'),'COD':p.get('COD'),'TSS':p.get('TSS'),
                'FCB':p.get('FCB'),'TCB':p.get('TCB'),'score':score,'status':score_to_overall_status(score)}

    for nm, tp, key in [('ก่อนบำบัด (บ่อรวบรวม)','ก่อนบำบัด','equalization'),
                         ('กำลังบำบัด (เติมอากาศ)','กำลังบำบัด','aeration'),
                         ('กำลังบำบัด (ตกตะกอน)','กำลังบำบัด','sedimentation'),
                         ('หลังบำบัด (บ่อพักน้ำทิ้ง)','หลังบำบัด','treated')]:
        r = stage_row_d(nm, tp, stage_latest.get(key))
        if r: stage_compare.append(r)

    if stage_latest.get('equalization') and stage_latest.get('treated'):
        b, a = build_parameters(stage_latest['equalization']), build_parameters(stage_latest['treated'])
        reduction = {k: percent_reduction(b.get(k), a.get(k)) for k in ['BOD','COD','TSS','FCB','TCB']}

    score_value, overall_status, overall_message = None, 'No Data', ''
    fuzzy_membership, parameter_results, param_rows = {}, {}, []
    compliance_counts = {'pass':0,'near':0,'fail':0,'na':0}
    reasons_fail, reasons_near, advanced_analysis = [], [], {}

    if after_data is not None:
        params = build_parameters(after_data)
        if can_fuzzy_eval(params):
            fuzzy_result = get_fuzzy_system(getattr(after_data,'building_type',None) or DEFAULT_BUILDING_TYPE).evaluate_overall(params)
        else:
            fuzzy_result = {'overall_score':None,'overall_level':'No Data',
                            'overall_message':'ข้อมูลไม่ครบสำหรับประเมิน Fuzzy',
                            'parameter_results':{},'fuzzy_membership':{}}
        score_value       = fuzzy_result.get('overall_score')
        overall_status    = score_to_overall_status(score_value)
        overall_message   = fuzzy_result.get('overall_message') or ''
        fuzzy_membership  = fuzzy_result.get('fuzzy_membership') or {}
        parameter_results = fuzzy_result.get('parameter_results') or {}
        advanced_analysis = advanced_water_analysis(parameter_results, params)

        for k in ['pH','BOD','COD','TSS','TDS','O&G','TKN','Sulfide','TCB','FCB','Cl2']:
            v   = params.get(k)
            std = find_std_for_param(k, getattr(after_data,'building_type',None))
            comp = classify_compliance(v, std, tol=0.10)
            if comp == 'Pass': compliance_counts['pass'] += 1
            elif comp == 'Near Limit': compliance_counts['near'] += 1; reasons_near.append(k)
            elif comp == 'Fail': compliance_counts['fail'] += 1; reasons_fail.append(k)
            else: compliance_counts['na'] += 1
            pr = parameter_results.get(k, {}) if isinstance(parameter_results, dict) else {}
            # ใช้ status จาก fuzzy ก่อน (รู้จัก "ไม่กำหนด") แล้ว fallback ไป classify_compliance
            fuzzy_status = pr.get('status') if isinstance(pr, dict) else None
            effective_status = fuzzy_status if fuzzy_status in ('Pass','Near Limit','Fail','ไม่กำหนด') else comp
            # risk_percent: None ถ้าไม่กำหนด, ค่าจริงถ้า regulated
            risk_raw = pr.get('risk_percent') if isinstance(pr,dict) else None
            risk_safe = risk_raw if isinstance(risk_raw, (int, float)) else 0
            param_rows.append({'label':k,'value':v,
                'standard':standard_text(std if isinstance(std,dict) else None),
                'status':effective_status,'fuzzy_term':pr.get('dominant_term'),
                'membership_degree':pr.get('dominant_mu'),
                'risk_percent':risk_raw,        # คงไว้เป็น None ถ้าไม่กำหนด — template เช็คเอง
                'risk_percent_safe':risk_safe,  # สำหรับ sort/compare
                'regulated': pr.get('regulated', True) if isinstance(pr,dict) else True,
                'risk_level':pr.get('status')})

    total_params    = len(param_rows)
    compliance_rate = round(compliance_counts['pass']/total_params*100,1) if total_params > 0 else 0.0
    param_stats     = calculate_detailed_statistics(selected_location, selected_pond, date_from, date_to)

    trend_labels, trend_bod, trend_cod, trend_tss, trend_ph, trend_coliform = [], [], [], [], [], []
    _show_trend_d = (date_from_str != date_to_str)
    if _show_trend_d:
        _latest_in_filter2 = WaterQualityData.query.filter(
            WaterQualityData.sample_location == selected_location,
            WaterQualityData.pond_name == selected_pond,
            WaterQualityData.sample_date.between(date_from, date_to)
        ).order_by(WaterQualityData.sample_date.desc()).first()
        if _latest_in_filter2:
            _trend_anchor2 = _latest_in_filter2.sample_date
            seven_days_ago = _trend_anchor2 - timedelta(days=7)
            _trend_raw = WaterQualityData.query.filter(
                WaterQualityData.sample_location == selected_location,
                WaterQualityData.pond_name == selected_pond,
                WaterQualityData.sample_date.between(seven_days_ago, _trend_anchor2)
            ).order_by(WaterQualityData.sample_date.asc()).all()
            _trend_by_date = defaultdict(list)
            for _r in _trend_raw: _trend_by_date[_r.sample_date.strftime('%d/%m/%Y')].append(_r)
            for _d in sorted(_trend_by_date.keys()):
                _rows = _trend_by_date[_d]
                def _avg2(attr, rows=_rows):
                    vals = [safe_float(getattr(r,attr)) for r in rows if safe_float(getattr(r,attr)) is not None]
                    return round(sum(vals)/len(vals),2) if vals else None
                trend_labels.append(_d[:-5])
                trend_bod.append(_avg2('BOD')); trend_cod.append(_avg2('COD'))
                trend_tss.append(_avg2('TSS')); trend_ph.append(_avg2('pH'))
                trend_coliform.append(_avg2('FCB'))
            if len(trend_labels) < 2:
                trend_labels, trend_bod, trend_cod, trend_tss, trend_ph, trend_coliform = [], [], [], [], [], []

    _all_stat_d = WaterQualityData.query.filter(
        WaterQualityData.sample_location == selected_location,
        WaterQualityData.pond_name == selected_pond,
        WaterQualityData.sample_date.between(date_from, date_to)
    ).all()
    def _stat_d(attr):
        vals = [safe_float(getattr(r,attr)) for r in _all_stat_d if safe_float(getattr(r,attr)) is not None]
        if not vals: return None, None, None
        return round(sum(vals)/len(vals),1), round(max(vals),1), round(min(vals),1)
    bod_mean, bod_max_s, bod_min = _stat_d('BOD')
    cod_mean, cod_max_s, cod_min = _stat_d('COD')
    tss_mean, tss_max_s, tss_min = _stat_d('TSS')

    return render_template('staff_home.html',
        location_list=location_list, pond_list=pond_list,
        selected_location=selected_location, selected_pond=selected_pond,
        display_date_text=display_date_text, date_mode=date_mode,
        date_from=date_from_str, date_to=date_to_str, selected_month=selected_month,
        mode=mode, after_data=after_data, after_count=after_count,
        pond_data=pond_data, score_value=score_value, overall_status=overall_status,
        overall_message=overall_message, fuzzy_membership=fuzzy_membership,
        parameter_results=parameter_results, param_rows=param_rows,
        compliance_counts=compliance_counts, compliance_rate=compliance_rate,
        reasons_fail=reasons_fail, reasons_near=reasons_near,
        reduction=reduction, stage_compare=stage_compare,
        param_stats=param_stats, water_standards=to_object(WATER_STANDARDS),
        advanced_analysis=advanced_analysis,
        trend_labels=trend_labels, trend_bod=trend_bod, trend_cod=trend_cod,
        trend_tss=trend_tss, trend_ph=trend_ph, trend_coliform=trend_coliform,
        bod_mean=bod_mean, bod_max=bod_max_s, bod_min=bod_min,
        cod_mean=cod_mean, cod_max=cod_max_s, cod_min=cod_min,
        tss_mean=tss_mean, tss_max=tss_max_s, tss_min=tss_min,
        available_dates=available_dates, snapshot_date=snapshot_date,
        recent_days=recent_days,
    )

# ⚠️ ลบออกหลังใช้งานแล้ว!
@app.route('/admin/clear-all-data', methods=['GET','POST'])
@login_required
def clear_all_data():
    if request.method == 'POST':
        try:
            EvaluationResult.query.delete()
            WaterQualityData.query.delete()
            db.session.commit()
            return '<h2 style="color:green">✅ ลบข้อมูลทั้งหมดแล้ว</h2><a href="/dashboard">กลับหน้าหลัก</a>'
        except Exception as e:
            db.session.rollback()
            return f'<h2 style="color:red">❌ Error: {e}</h2>'
    return '''
        <h2>⚠️ ยืนยันลบข้อมูลทั้งหมด?</h2>
        <p style="color:red">การกระทำนี้ไม่สามารถย้อนกลับได้</p>
        <form method="POST">
            <button type="submit" style="background:red;color:white;padding:10px 20px;font-size:16px;border:none;cursor:pointer;border-radius:6px">
                ยืนยัน ลบทั้งหมด
            </button>
            <a href="/dashboard" style="margin-left:12px">ยกเลิก</a>
        </form>
    '''
# =========================
# API: list distinct locations & ponds
# =========================
@app.route('/api/locations-and-ponds')
@login_required
def api_locations_and_ponds():
    """คืน dict: { 'locations': [...], 'ponds_by_location': { 'รพ. A': ['บ่อ1','บ่อ2'], ... } }
    เพื่อใช้กับ dropdown ใน add_data / upload_data — กันพิมพ์ผิด/ขาดช่อง"""
    locations = sorted({
        (r[0] or '').strip()
        for r in db.session.query(WaterQualityData.sample_location).distinct().all()
        if r[0] and r[0].strip()
    })
    ponds_by_loc = {}
    rows = db.session.query(
        WaterQualityData.sample_location, WaterQualityData.pond_name
    ).distinct().all()
    for loc, pond in rows:
        loc  = (loc  or '').strip()
        pond = (pond or '').strip()
        if not loc or not pond: continue
        ponds_by_loc.setdefault(loc, set()).add(pond)
    ponds_by_loc = {k: sorted(v) for k, v in ponds_by_loc.items()}
    return jsonify({
        'locations': locations,
        'ponds_by_location': ponds_by_loc,
    })


# =========================
# REFERENCE TABLE: ช่วงค่า μ=1 (full score) per building type
# =========================
@app.route('/reference/zones')
@login_required
def reference_zones():
    """หน้าตารางอ้างอิง: ช่วงค่าที่ได้คะแนนเต็ม (μ=1) ของแต่ละพารามิเตอร์
    ในแต่ละประเภทอาคาร ก/ข/ค/ง"""
    selected_bt = (request.args.get('bt') or DEFAULT_BUILDING_TYPE).strip()
    if selected_bt not in BUILDING_TYPES:
        selected_bt = DEFAULT_BUILDING_TYPE

    # คำนวณ zones ของทุกประเภทอาคาร (cache อยู่แล้วใน fuzzy_logic)
    zones_by_type = {}
    for bt in BUILDING_TYPES:
        zones_by_type[bt] = get_param_zones(bt)

    return render_template('reference_zones.html',
                           selected_bt=selected_bt,
                           building_types=BUILDING_TYPES,
                           zones_by_type=zones_by_type)


# =========================
# ADD DATA
# =========================
@app.route('/data/add', methods=['GET','POST'])
@login_required
def add_data():
    if request.method == 'POST':
        try:
            sample_location = request.form.get('sample_location','').strip()
            if not sample_location: flash('กรุณาระบุสถานที่เก็บตัวอย่าง','error'); return redirect(url_for('add_data'))
            sample_date_raw = request.form.get('sample_date')
            if not sample_date_raw: flash('กรุณาระบุวันเวลาเก็บตัวอย่าง','error'); return redirect(url_for('add_data'))
            sample_date = datetime.strptime(sample_date_raw,'%Y-%m-%dT%H:%M')
            pond_name   = request.form.get('pond_name','').strip()
            if not pond_name: flash('กรุณาระบุชื่อบ่อ/ระบบบำบัด','error'); return redirect(url_for('add_data'))
            sample_type = request.form.get('sample_type','หลังบำบัด')
            building_type = request.form.get('building_type', DEFAULT_BUILDING_TYPE).strip() or DEFAULT_BUILDING_TYPE
            if building_type not in BUILDING_TYPES:
                building_type = DEFAULT_BUILDING_TYPE
            water_data  = WaterQualityData(
                user_id=current_user.id, sample_date=sample_date,
                sample_location=sample_location, pond_name=pond_name, sample_type=sample_type,
                building_type=building_type,
                pH=get_float(request.form,'pH'), BOD=get_float(request.form,'BOD'),
                COD=get_float(request.form,'COD'), TSS=get_float(request.form,'TSS'),
                TDS=get_float(request.form,'TDS'), oil_grease=get_float(request.form,'oil_grease'),
                TKN=get_float(request.form,'TKN'), sulfide=get_float(request.form,'sulfide'),
                TCB=get_float(request.form,'TCB'), FCB=get_float(request.form,'FCB'),
                chlorine=get_float(request.form,'chlorine'),
                ammonium=get_float(request.form,'ammonium'),
                dissolved_oxygen=get_float(request.form,'dissolved_oxygen'),
            )
            db.session.add(water_data); db.session.flush()
            if sample_type == 'หลังบำบัด':
                params = build_parameters(water_data)
                if can_fuzzy_eval(params):
                    evaluation = get_fuzzy_system(building_type).evaluate_overall(params)
                    score = evaluation.get('overall_score'); pr = evaluation.get('parameter_results',{}) or {}
                    result = EvaluationResult(
                        water_data_id=water_data.id, overall_status=score_to_overall_status(score),
                        overall_score=score, overall_message=evaluation.get('overall_message',''),
                        pass_count=0, near_limit_count=0, fail_count=0, total_parameters=0,
                        pH_result=pr.get('pH'), BOD_result=pr.get('BOD'), COD_result=pr.get('COD'),
                        TSS_result=pr.get('TSS'), TDS_result=pr.get('TDS'),
                        oil_grease_result=pr.get('O&G'), TKN_result=pr.get('TKN'),
                        sulfide_result=pr.get('Sulfide'), TCB_result=pr.get('TCB'),
                        FCB_result=pr.get('FCB'), chlorine_result=pr.get('Cl2'),
                    )
                    db.session.add(result); db.session.commit()
                    flash('บันทึกข้อมูลและประเมินเรียบร้อย!','success')
                    return redirect(url_for('view_result', result_id=result.id))
                else:
                    db.session.commit()
                    flash('บันทึกข้อมูลแล้ว แต่ข้อมูลไม่ครบสำหรับประเมิน Fuzzy','warning')
                    return redirect(url_for('dashboard'))
            db.session.commit(); flash('บันทึกข้อมูลเรียบร้อย!','success')
            return redirect(url_for('dashboard'))
        except Exception as e:
            db.session.rollback(); flash(f'เกิดข้อผิดพลาด: {str(e)}','error')
            return redirect(url_for('add_data'))
    return render_template('add_data.html',
                           water_standards=to_object(WATER_STANDARDS),
                           water_standards_by_type={bt: to_object(WATER_STANDARDS_BY_TYPE[bt]) for bt in BUILDING_TYPES},
                           building_types=BUILDING_TYPES,
                           default_building_type=DEFAULT_BUILDING_TYPE)

# =========================
# UPLOAD DATA
# =========================
@app.route('/data/upload', methods=['GET','POST'])
@login_required
def upload_data():
    if request.method != 'POST':
        return render_template('upload_data.html',
                               water_standards=to_object(WATER_STANDARDS),
                               building_types=BUILDING_TYPES,
                               default_building_type=DEFAULT_BUILDING_TYPE)
    if 'file' not in request.files:
        flash('ไม่พบไฟล์','error'); return redirect(request.url)
    file = request.files['file']
    if file.filename == '':
        flash('ไม่ได้เลือกไฟล์','error'); return redirect(request.url)
    if not (file and allowed_file(file.filename)):
        flash('ไฟล์ต้องเป็น .xlsx หรือ .xls เท่านั้น','error'); return redirect(request.url)

    # default building_type จากฟอร์ม (ใช้กับแถวที่ไม่ระบุในไฟล์ Excel)
    default_bt = (request.form.get('building_type') or DEFAULT_BUILDING_TYPE).strip() or DEFAULT_BUILDING_TYPE
    if default_bt not in BUILDING_TYPES:
        default_bt = DEFAULT_BUILDING_TYPE

    try:
        # อ่านแบบ format-aware (รองรับทั้งรูปแบบ flat เดิม + รูปแบบโรงพยาบาล)
        # พร้อมแปลง 'ไม่พบ'→0, 'Ozone'→Cl2=1.0 ฯลฯ และข้ามแถว 'มาตรฐาน'/'หมายเหตุ'
        df, subs_notes, fmt = normalize_hospital_xlsx(file, fuzzy_system=fuzzy_system)

        if df is None or len(df) == 0:
            flash('ไฟล์ไม่มีแถวข้อมูลที่อ่านได้','error')
            return redirect(request.url)

        # สร้าง map ชื่อคอลัมน์แบบ case-insensitive
        cols_lower = {str(c).strip().lower(): c for c in df.columns}
        def gv(row, *names):
            for n in names:
                key = cols_lower.get(n.lower())
                if key is not None:
                    v = row.get(key)
                    if v is not None:
                        return v
            return None

        # ตรวจสอบคอลัมน์ขั้นต่ำ
        required_min = ['sample_date','sample_location','ph','bod','cod',
                        'tss','tds','o&g','tkn','sulfide','tcb','fcb','cl2']
        missing = [c for c in required_min if c not in cols_lower]
        if missing:
            flash(f'ไฟล์ Excel ขาดคอลัมน์: {", ".join(missing)} (รูปแบบที่ตรวจพบ: {fmt})','error')
            return redirect(request.url)

        success_count = error_count = 0
        valid_dates = []
        last_location = None

        for _, row in df.iterrows():
            try:
                with db.session.begin_nested():
                    loc  = gv(row, 'sample_location')
                    pond = gv(row, 'pond_name') or 'บ่อพักน้ำทิ้ง'
                    if is_missing(loc):
                        error_count += 1; continue
                    sample_location = str(loc).strip()
                    pond_name       = str(pond).strip()

                    st = gv(row, 'sample_type')
                    sample_type = 'หลังบำบัด' if is_missing(st) else str(st).strip()

                    # building_type: รองรับคอลัมน์ 'building_type' หรือ 'ประเภทอาคาร' ในไฟล์
                    bt_raw = gv(row, 'building_type', 'ประเภทอาคาร', 'อาคาร')
                    row_bt = default_bt if is_missing(bt_raw) else str(bt_raw).strip()
                    # ถ้าระบุเป็น "อาคารประเภท ก" ให้ตัดคำว่า อาคารประเภท ออก
                    row_bt = row_bt.replace('อาคารประเภท','').replace('ประเภทอาคาร','').replace('อาคาร','').strip()
                    if row_bt not in BUILDING_TYPES:
                        row_bt = default_bt

                    try:
                        sample_date = pd.to_datetime(gv(row, 'sample_date'))
                        if pd.isna(sample_date): raise ValueError
                    except:
                        error_count += 1; continue

                    existing = WaterQualityData.query.filter_by(
                        sample_date=sample_date, sample_location=sample_location,
                        pond_name=pond_name, sample_type=sample_type).first()
                    if existing:
                        water_data = existing
                    else:
                        water_data = WaterQualityData(
                            user_id=current_user.id, sample_date=sample_date,
                            sample_location=sample_location,
                            pond_name=pond_name, sample_type=sample_type,
                            building_type=row_bt)
                        db.session.add(water_data)

                    water_data.user_id     = current_user.id
                    water_data.building_type = row_bt
                    water_data.pH          = safe_float(gv(row, 'pH'))
                    water_data.BOD         = safe_float(gv(row, 'BOD'))
                    water_data.COD         = safe_float(gv(row, 'COD'))
                    water_data.TSS         = safe_float(gv(row, 'TSS'))
                    water_data.TDS         = safe_float(gv(row, 'TDS'))
                    water_data.oil_grease  = safe_float(gv(row, 'O&G'))
                    water_data.TKN         = safe_float(gv(row, 'TKN'))
                    water_data.sulfide     = safe_float(gv(row, 'Sulfide','sulfide'))
                    water_data.TCB         = safe_float(gv(row, 'TCB'))
                    water_data.FCB         = safe_float(gv(row, 'FCB'))
                    water_data.chlorine    = safe_float(gv(row, 'Cl2'))
                    water_data.ammonium    = safe_float(gv(row, 'NH4'))
                    water_data.dissolved_oxygen = safe_float(gv(row, 'DO'))

                    db.session.flush()
                    if water_data.sample_type == 'หลังบำบัด':
                        params = build_parameters(water_data)
                        if can_fuzzy_eval(params):
                            evaluation = get_fuzzy_system(water_data.building_type).evaluate_overall(params)
                            pr = evaluation.get('parameter_results',{}) or {}
                            overall_score = evaluation.get('overall_score')
                            if overall_score is not None:
                                result = EvaluationResult.query.filter_by(water_data_id=water_data.id).first()
                                if not result:
                                    result = EvaluationResult(water_data_id=water_data.id); db.session.add(result)
                                result.overall_status  = score_to_overall_status(overall_score)
                                result.overall_score   = float(overall_score)
                                result.overall_message = evaluation.get('overall_message','')
                                result.pass_count = result.near_limit_count = result.fail_count = result.total_parameters = 0
                                result.pH_result=pr.get('pH'); result.BOD_result=pr.get('BOD')
                                result.COD_result=pr.get('COD'); result.TSS_result=pr.get('TSS')
                                result.TDS_result=pr.get('TDS'); result.oil_grease_result=pr.get('O&G')
                                result.TKN_result=pr.get('TKN'); result.sulfide_result=pr.get('Sulfide')
                                result.TCB_result=pr.get('TCB'); result.FCB_result=pr.get('FCB')
                                result.chlorine_result=pr.get('Cl2')
                    success_count += 1
                    valid_dates.append(sample_date)
                    last_location = sample_location
            except IntegrityError:
                error_count += 1; continue
            except Exception as ex:
                app.logger.error(f'[upload row] {ex}')
                error_count += 1; continue

        db.session.commit()

        # สรุปข้อความ + หมายเหตุการแปลงค่า
        fmt_label = 'รูปแบบโรงพยาบาล (หลายชีต)' if fmt == 'hospital' else 'รูปแบบตารางเดิม'
        msg = f'อัปโหลดสำเร็จ {success_count} รายการ ({fmt_label})'
        if error_count > 0:
            msg += f' · ข้าม {error_count} รายการ'
        if subs_notes:
            n_ozone = sum(1 for n in subs_notes if 'Ozone' in str(n.get('note','')))
            n_nf    = sum(1 for n in subs_notes if 'ไม่พบ' in str(n.get('note','')))
            parts = []
            if n_nf:    parts.append(f'"ไม่พบ"→0 ({n_nf} เซลล์)')
            if n_ozone: parts.append(f'"Ozone"→Cl2 {CL2_OZONE_SUBSTITUTE} mg/L ({n_ozone} เซลล์)')
            if parts:
                msg += ' · แทนค่า: ' + ' , '.join(parts)
        flash(msg,'success')

        latest_date_str = max(valid_dates).strftime('%Y-%m-%d') if valid_dates else None
        if latest_date_str and last_location:
            return redirect(url_for('dashboard',location=last_location,date_from=latest_date_str,date_to=latest_date_str))
        return redirect(url_for('dashboard'))
    except Exception as e:
        import traceback
        db.session.rollback()
        app.logger.error(f'[upload] {traceback.format_exc()}')
        flash(f'เกิดข้อผิดพลาดในการอ่านไฟล์: {type(e).__name__}: {str(e)}','error')
        return redirect(request.url)

@app.route('/data/edit/<int:water_data_id>', methods=['GET','POST'])
@login_required
def edit_data(water_data_id):
    water_data = WaterQualityData.query.get_or_404(water_data_id)
    if request.method == 'POST':
        try:
            sample_location = (request.form.get('sample_location') or '').strip()
            pond_name       = (request.form.get('pond_name') or '').strip()
            sample_type     = (request.form.get('sample_type') or '').strip() or 'หลังบำบัด'
            building_type   = (request.form.get('building_type') or DEFAULT_BUILDING_TYPE).strip() or DEFAULT_BUILDING_TYPE
            if building_type not in BUILDING_TYPES:
                building_type = DEFAULT_BUILDING_TYPE
            sample_date_raw = request.form.get('sample_date')
            if not sample_date_raw: flash('กรุณาระบุวันเวลาเก็บตัวอย่าง','error'); return redirect(url_for('edit_data',water_data_id=water_data_id))
            sample_date = datetime.strptime(sample_date_raw,'%Y-%m-%dT%H:%M')
            if not sample_location: flash('กรุณาระบุสถานที่','error'); return redirect(url_for('edit_data',water_data_id=water_data_id))
            if not pond_name: flash('กรุณาระบุชื่อบ่อ','error'); return redirect(url_for('edit_data',water_data_id=water_data_id))
            water_data.sample_location=sample_location; water_data.pond_name=pond_name
            water_data.sample_type=sample_type; water_data.sample_date=sample_date
            water_data.building_type=building_type
            water_data.pH=get_float(request.form,'pH'); water_data.BOD=get_float(request.form,'BOD')
            water_data.COD=get_float(request.form,'COD'); water_data.TSS=get_float(request.form,'TSS')
            water_data.TDS=get_float(request.form,'TDS'); water_data.oil_grease=get_float(request.form,'oil_grease')
            water_data.TKN=get_float(request.form,'TKN'); water_data.sulfide=get_float(request.form,'sulfide')
            water_data.TCB=get_float(request.form,'TCB'); water_data.FCB=get_float(request.form,'FCB')
            water_data.chlorine=get_float(request.form,'chlorine')
            water_data.ammonium=get_float(request.form,'ammonium')
            water_data.dissolved_oxygen=get_float(request.form,'dissolved_oxygen')
            db.session.flush()
            if water_data.sample_type == 'หลังบำบัด':
                params     = build_parameters(water_data)
                evaluation = (get_fuzzy_system(water_data.building_type).evaluate_overall(params) if can_fuzzy_eval(params)
                              else {'overall_score':None,'overall_message':'ข้อมูลไม่ครบ','parameter_results':{}})
                score  = evaluation.get('overall_score')
                result = EvaluationResult.query.filter_by(water_data_id=water_data.id).first()
                if not result: result = EvaluationResult(water_data_id=water_data.id); db.session.add(result)
                result.overall_status  = score_to_overall_status(score)
                result.overall_score   = score
                result.overall_message = evaluation.get('overall_message','')
                pr = evaluation.get('parameter_results',{}) or {}
                result.pH_result=pr.get('pH'); result.BOD_result=pr.get('BOD')
                result.COD_result=pr.get('COD'); result.TSS_result=pr.get('TSS')
                result.TDS_result=pr.get('TDS'); result.oil_grease_result=pr.get('O&G')
                result.TKN_result=pr.get('TKN'); result.sulfide_result=pr.get('Sulfide')
                result.TCB_result=pr.get('TCB'); result.FCB_result=pr.get('FCB')
                result.chlorine_result=pr.get('Cl2')
            db.session.commit(); flash('บันทึกการแก้ไขเรียบร้อย!','success')
            return redirect(url_for('dashboard',location=water_data.sample_location,pond=water_data.pond_name))
        except Exception as e:
            db.session.rollback(); flash(f'เกิดข้อผิดพลาด: {str(e)}','error')
            return redirect(url_for('edit_data',water_data_id=water_data_id))
    return render_template('edit.html', water_data=water_data,
                           water_standards=to_object(get_water_standards(water_data.building_type)),
                           water_standards_by_type={bt: to_object(WATER_STANDARDS_BY_TYPE[bt]) for bt in BUILDING_TYPES},
                           building_types=BUILDING_TYPES,
                           default_building_type=DEFAULT_BUILDING_TYPE)

# ── delete_data ── แก้ตรง permission check ──────────────────
@app.route('/data/delete/<int:water_data_id>', methods=['POST'])
@login_required
def delete_data(water_data_id):
    data = WaterQualityData.query.get_or_404(water_data_id)

    # ✅ แก้: อนุญาตให้ลบได้ถ้าเป็นเจ้าของ หรือเป็น admin
    # ถ้าต้องการให้ทุก user ลบได้ทุกรายการ → ลบ if block นี้ออกทั้งหมด
    is_owner = (getattr(data, 'user_id', None) == current_user.id)
    is_admin = getattr(current_user, 'is_admin', False)
    if not (is_owner or is_admin):
        flash('คุณไม่มีสิทธิ์ลบข้อมูลนี้', 'error')
        return redirect(url_for('history'))

    try:
        result = EvaluationResult.query.filter_by(water_data_id=data.id).first()
        if result:
            db.session.delete(result)
        db.session.delete(data)
        db.session.commit()
        flash('ลบข้อมูลเรียบร้อยแล้ว', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'ลบไม่สำเร็จ: {str(e)}', 'error')

    return redirect(url_for('history', page=request.args.get('page', 1, type=int)))
@app.route('/result/<int:result_id>')
@login_required
def view_result(result_id):
    result     = EvaluationResult.query.get_or_404(result_id)
    water_data = result.water_data
    bt = getattr(water_data, 'building_type', None) or DEFAULT_BUILDING_TYPE
    parameter_results = []; live_pr = {}
    try:
        params_live = build_parameters(water_data)
        if can_fuzzy_eval(params_live):
            live_pr = get_fuzzy_system(bt).evaluate_overall(params_live).get('parameter_results',{}) or {}
    except: pass
    for param_name, param_result in [
        ('pH',result.pH_result),('BOD',result.BOD_result),('COD',result.COD_result),
        ('TSS',result.TSS_result),('TDS',result.TDS_result),('O&G',result.oil_grease_result),
        ('TKN',result.TKN_result),('Sulfide',result.sulfide_result),
        ('TCB',result.TCB_result),('FCB',result.FCB_result),('Cl2',result.chlorine_result),
    ]:
        if param_result:
            std      = find_std_for_param(param_name, bt) or {}
            enriched = dict(param_result) if isinstance(param_result,dict) else {}
            enriched['unit']       = std.get('unit','')
            enriched['standard']   = standard_text(std)
            enriched['regulated']  = std.get('regulated', True) if std else True
            if 'risk_percent' not in enriched:
                enriched['risk_percent'] = (live_pr.get(param_name,{}) or {}).get('risk_percent',0)
            if 'status' not in enriched:
                enriched['status'] = (live_pr.get(param_name,{}) or {}).get('status','N/A')
            parameter_results.append({'name':std.get('name',param_name),'data':enriched})
    return render_template('result.html', result=result, water_data=water_data,
                           parameter_results=parameter_results,
                           water_standards=to_object(get_water_standards(bt)),
                           building_type=bt,
                           wqi_bands=WQI_BANDS)

# =========================
# HISTORY
# =========================
@app.route('/history')
@login_required
def history():
    page        = request.args.get('page', 1, type=int)
    q           = request.args.get('q', '').strip()
    status      = request.args.get('status', '')
    location    = request.args.get('location', '')
    sample_type = request.args.get('sample_type', '')

    query = WaterQualityData.query
    if q:
        like = f'%{q}%'
        query = query.filter(db.or_(
            WaterQualityData.sample_location.ilike(like),
            WaterQualityData.pond_name.ilike(like),
            WaterQualityData.sample_type.ilike(like),
        ))
    if location:   query = query.filter(WaterQualityData.sample_location == location)
    if sample_type: query = query.filter(WaterQualityData.sample_type == sample_type)

    if status == 'pass':
        query = query.join(EvaluationResult, EvaluationResult.water_data_id == WaterQualityData.id).filter(
            EvaluationResult.overall_status == 'Pass')
    elif status == 'near':
        query = query.join(EvaluationResult, EvaluationResult.water_data_id == WaterQualityData.id).filter(
            EvaluationResult.overall_status == 'Near Limit')
    elif status == 'fail':
        query = query.join(EvaluationResult, EvaluationResult.water_data_id == WaterQualityData.id).filter(
            EvaluationResult.overall_status == 'Fail')
    elif status == 'skip':
        query = query.filter(db.or_(
            WaterQualityData.sample_type != 'หลังบำบัด',
            ~WaterQualityData.evaluation.has()
        ))

    query      = query.order_by(WaterQualityData.id.desc())
    pagination = query.paginate(page=page, per_page=app.config['ITEMS_PER_PAGE'], error_out=False)
    all_locations = [
        r[0] for r in db.session.query(WaterQualityData.sample_location)
        .distinct()
        .filter(WaterQualityData.sample_location.isnot(None), WaterQualityData.sample_location != '')
        .order_by(WaterQualityData.sample_location)
        .all()
    ]
    return render_template('history.html', pagination=pagination,
                           water_standards=to_object(WATER_STANDARDS), all_locations=all_locations)


# =========================
# HELPER: Error page for export failures
# =========================
def _err_page(title, detail, hint=''):
    return (f'<!DOCTYPE html><html lang="th"><head><meta charset="UTF-8">'
            f'<title>Export Error</title>'
            f'<style>body{{font-family:sans-serif;padding:40px;background:#f8fafc}}'
            f'h2{{color:#dc2626}}pre{{background:#fee2e2;padding:14px;border-radius:8px;'
            f'font-size:13px;white-space:pre-wrap}}</style></head><body>'
            f'<h2>&#9888; Export ไม่สำเร็จ</h2><p><strong>{title}</strong></p>'
            f'<pre>{detail}</pre>'
            f'{("<p>"+hint+"</p>") if hint else ""}'
            f'<p><a href="javascript:window.close()">ปิด</a></p></body></html>')


# =========================
# EXPORT: Monthly Report (Excel — รูปแบบต่อโรงพยาบาล + ชีตกราฟ)
# =========================
@app.route('/export/monthly-report')
@login_required
def export_monthly_report():
    """สร้างไฟล์ Excel รายเดือน — มี progressive fallback เพื่อไม่ให้
    ผู้ใช้ติดกับดัก "ไม่มีข้อมูล → redirect กลับ dashboard"

    URL params (ทั้งหมดเป็น optional):
      - location   : ชื่อโรงพยาบาล (ถ้าว่าง = ทุกโรงพยาบาล)
      - month      : YYYY-MM  (ถ้าว่าง/'all' = ทุกเดือน)
      - sample_type: default 'หลังบำบัด' (ถ้าระบุ 'all' = ทุก type)

    ลำดับ fallback:
      1) กรองครบทุกเงื่อนไข
      2) ถ้าว่าง: ลบ month filter
      3) ถ้ายังว่าง: ลบ sample_type filter
      4) ถ้ายังว่าง: สร้างไฟล์ Excel พร้อม sheet "ไม่พบข้อมูล" (ยัง download ได้)
    """
    location    = (request.args.get('location') or '').strip()
    month_str   = (request.args.get('month') or '').strip()
    sample_type = (request.args.get('sample_type') or 'หลังบำบัด').strip()

    # ── Parse month ──────────────────────────────────────────────
    dfrom = dto = None
    y = m = None
    if month_str and month_str.lower() not in ('all', 'none', 'any'):
        try:
            y, m = map(int, month_str.split('-'))
            dfrom = datetime(y, m, 1)
            dto   = (datetime(y+1, 1, 1) - timedelta(seconds=1) if m == 12
                     else datetime(y, m+1, 1) - timedelta(seconds=1))
        except Exception:
            app.logger.warning(f'[export_monthly] invalid month={month_str!r}, ignore filter')
            dfrom = dto = None
            month_str = 'all'

    def _query(use_month=True, use_type=True):
        q = WaterQualityData.query
        if location:
            q = q.filter(WaterQualityData.sample_location == location)
        if use_type and sample_type and sample_type.lower() != 'all':
            q = q.filter(WaterQualityData.sample_type == sample_type)
        if use_month and dfrom and dto:
            q = q.filter(WaterQualityData.sample_date.between(dfrom, dto))
        return q.order_by(WaterQualityData.sample_location.asc(),
                          WaterQualityData.sample_date.asc()).all()

    try:
        # ── Progressive fallback ─────────────────────────────────
        fallback_note = ''
        all_rows = _query(use_month=True, use_type=True)
        if not all_rows and dfrom:
            all_rows = _query(use_month=False, use_type=True)
            if all_rows:
                fallback_note = (f'ไม่พบข้อมูลในเดือน {month_str} '
                                 f'→ แสดงข้อมูลทุกช่วงเวลาแทน')
        if not all_rows:
            all_rows = _query(use_month=False, use_type=False)
            if all_rows:
                fallback_note = (f'ไม่พบข้อมูลประเภท "{sample_type}" → '
                                 f'แสดงทุกประเภทตัวอย่าง')

        app.logger.info(
            f'[export_monthly] location={location!r} month={month_str!r} '
            f'type={sample_type!r} → {len(all_rows)} rows '
            f'(fallback: {fallback_note or "none"})'
        )

        # ── จัดกลุ่มตามสถานที่ ──────────────────────────────────
        hosp_rows = {}
        for r in all_rows:
            hosp_rows.setdefault(r.sample_location or '–', []).append(r)

        # ── Title suffix ──────────────────────────────────────
        TH_M = {1:'มกราคม',2:'กุมภาพันธ์',3:'มีนาคม',4:'เมษายน',
                5:'พฤษภาคม',6:'มิถุนายน',7:'กรกฎาคม',8:'สิงหาคม',
                9:'กันยายน',10:'ตุลาคม',11:'พฤศจิกายน',12:'ธันวาคม'}
        if y and m:
            title_suffix = f'ประจำเดือน {TH_M.get(m,str(m))} พ.ศ. {y+543}'
        else:
            title_suffix = 'ทุกช่วงเวลา'
        if fallback_note:
            title_suffix += f'  ({fallback_note})'

        # ── ถ้ายังไม่มีข้อมูลเลย: สร้างไฟล์ "No data" ─────────────
        if not hosp_rows:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook()
            ws = wb.active
            ws.title = 'ไม่พบข้อมูล'
            ws.merge_cells('A1:F1')
            ws['A1'] = 'ไม่พบข้อมูลในระบบ'
            ws['A1'].font = Font(name='Tahoma', bold=True, size=14, color='FFFFFF')
            ws['A1'].fill = PatternFill('solid', fgColor='DC2626')
            ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 28

            details = [
                ('เงื่อนไขที่ค้นหา', ''),
                ('  • โรงพยาบาล', location or '(ทุกโรงพยาบาล)'),
                ('  • เดือน', month_str or '(ทุกเดือน)'),
                ('  • ประเภทตัวอย่าง', sample_type),
                ('', ''),
                ('คำแนะนำ',
                 'กรุณาอัปโหลดข้อมูลผ่านหน้า "อัปโหลดข้อมูล" '
                 'หรือเพิ่มข้อมูลด้วยมือผ่านหน้า "เพิ่มข้อมูล" ก่อน'),
            ]
            for i, (k, v) in enumerate(details, start=3):
                ws.cell(row=i, column=1, value=k).font = Font(
                    name='Tahoma', bold=bool(k) and not k.startswith(' '),
                    size=11)
                ws.cell(row=i, column=2, value=v).font = Font(
                    name='Tahoma', size=11)
            ws.column_dimensions['A'].width = 26
            ws.column_dimensions['B'].width = 60

            buf = io.BytesIO()
            wb.save(buf); buf.seek(0)
            xlsx_bytes = buf.read()
        else:
            xlsx_bytes = build_hospital_xlsx_report(
                hosp_rows,
                fuzzy_system=fuzzy_system,
                get_fuzzy_system=get_fuzzy_system,
                default_building_type=DEFAULT_BUILDING_TYPE,
                build_parameters=build_parameters,
                can_fuzzy_eval=can_fuzzy_eval,
                score_to_overall_status=score_to_overall_status,
                standards=fuzzy_system.standards,
                standards_by_type={bt: BUILDING_STANDARDS[bt] for bt in BUILDING_TYPES},
                title_suffix=title_suffix,
            )

        # ── ชื่อไฟล์ (ไทย; helper จะจัดการ latin-1 safety) ────────
        parts = ['WQ_Monthly']
        parts.append(location if location else 'ALL')
        parts.append(month_str if month_str and month_str.lower() not in ('all','none','any')
                                else 'ALL-TIME')
        filename_thai = '_'.join(parts) + '.xlsx'

        return _xlsx_download_response(xlsx_bytes, filename_thai)
    except Exception as e:
        import traceback
        app.logger.error(f'[export_monthly_xlsx] {traceback.format_exc()}')
        return _err_page(f'เกิดข้อผิดพลาด: {type(e).__name__}',
                         traceback.format_exc()), 500


# =========================
# EXPORT: Single Result (Excel)
# =========================
@app.route('/export/result/<int:result_id>')
@login_required
def export_result_pdf(result_id):
    """ส่งออกผลการประเมินรายครั้งเป็นไฟล์ Excel

    URL path เดิมคงไว้เพื่อให้ templates ที่อ้างถึง endpoint นี้ยังใช้งานได้
    """
    try:
        result     = EvaluationResult.query.get_or_404(result_id)
        water_data = result.water_data
        username   = water_data.user.username if getattr(water_data,'user',None) else None
        sd         = water_data.sample_date

        _bt = getattr(water_data, 'building_type', None) or DEFAULT_BUILDING_TYPE
        _fs = get_fuzzy_system(_bt)
        # ใช้ WATER_STANDARDS_BY_TYPE (มีฟิลด์ regulated, display name) แทน fs.standards (raw dict)
        _stds = WATER_STANDARDS_BY_TYPE.get(_bt, WATER_STANDARDS_BY_TYPE[DEFAULT_BUILDING_TYPE])
        xlsx_bytes = build_single_result_xlsx(
            water_data,
            evaluation=result,
            fuzzy_system=_fs,
            build_parameters=build_parameters,
            can_fuzzy_eval=can_fuzzy_eval,
            score_to_overall_status=score_to_overall_status,
            standards=_stds,
            username=username,
            building_type=_bt,
        )

        ds = sd.strftime('%Y%m%d') if sd else 'nd'
        filename_thai = f'WQ_Result_{result_id}_{ds}.xlsx'
        return _xlsx_download_response(xlsx_bytes, filename_thai)
    except Exception as e:
        import traceback
        app.logger.error(f'[export_result_xlsx] id={result_id}: {traceback.format_exc()}')
        return _err_page(f'เกิดข้อผิดพลาด: {type(e).__name__}', traceback.format_exc()), 500





# =========================
# ERROR HANDLERS
# =========================
@app.errorhandler(404)
def not_found(error):
    return render_template('404.html', water_standards=to_object(WATER_STANDARDS)), 404

@app.errorhandler(500)
def internal_error(error):
    db.session.rollback()
    return render_template('500.html', water_standards=to_object(WATER_STANDARDS)), 500



# =========================
# STARTUP
# =========================
if __name__ == '__main__':
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    with app.app_context():
        db.create_all()
        check_and_sync_fuzzy()
    app.run(debug=True, host='0.0.0.0', port=5000)